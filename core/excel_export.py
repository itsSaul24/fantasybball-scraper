import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# (column, group, plain-English meaning, how to use it live)
COLUMN_GUIDE = [
    ("rank", "Identity", "Position in the pool, ranked by availability-adjusted value.", "Rough draft order. Not a bid price."),
    ("player", "Identity", "Player name.", ""),
    ("position", "Identity", "Primary ESPN position.", "Check against your unfilled slots."),
    ("eligible", "Identity", "Every roster slot ESPN lets this player fill.", "Multi-slot players add lineup flexibility; worth a small premium."),
    ("team", "Identity", "NBA team for the upcoming season.", ""),
    ("team_changed", "Identity", "TRUE if he changed teams this offseason.", "Last season's stats were produced somewhere else — discount them."),
    ("age", "Identity", "Age.", "Pair with career_arc: decline at 34 differs from decline at 24."),
    ("is_rookie", "Identity", "TRUE for incoming rookies with no NBA stats.", "Valued from community discussion only. Treat as a lottery ticket."),
    ("tier", "Verdict", "anchor / core / value / depth / dart.", "Anchors are roster centerpieces; darts are $1-3 fliers."),
    ("trend", "Verdict", "rising / falling / stable / volatile.", "Direction of travel into this season."),
    ("market_vs_value", "Verdict", "Where community sentiment sits vs actual production.", "'underrated' = bargain candidate. 'overhyped' = let the room overpay."),
    ("steal", "Pricing", "At or below this price, buy immediately.", "If bidding is at or under this, raise your hand."),
    ("target", "Pricing", "Realistic expected price; your plan number.", "Budget with this figure."),
    ("walk_away", "Pricing", "Stop bidding here. One more dollar is an overpay.", "Hard stop. Pivot to your alternative."),
    ("mechanical_value", "Pricing", "Formula price from last season's stats, before any judgement.", "A sanity anchor. Large gaps vs target mean the narrative moved the price."),
    ("community_read", "Analysis", "What Reddit actually says, synthesized across commenters.", "Says 'no meaningful discussion found' when the community is silent."),
    ("why", "Analysis", "The case for the price, reconciling community view with the numbers.", "Read before bidding."),
    ("risk", "Analysis", "The single most likely way this pick fails.", "Your pre-mortem."),
    ("live_note", "Analysis", "Tactical guidance for the auction room.", "When to push, when to fold, who to pivot to."),
    ("trajectory", "Analysis", "Year-over-year change, final-20-game form, and minutes trend.", "Late-season collapse plus falling minutes signals a shrinking role."),
    ("career_arc", "History", "Five-year shape: ascent, plateau, decline, with % of peak.", "Separates a rebound candidate from a fading veteran."),
    ("fp_trend_5y", "History", "Fantasy points per game gained/lost per season (fitted slope).", "Positive = climbing. Check seasons_of_data before trusting it."),
    ("fp_peak", "History", "Best single-season fantasy points per game in the last five.", "The realistic ceiling."),
    ("fp_peak_season", "History", "Which season the peak happened.", "A peak four years ago is a warning."),
    ("pct_of_peak", "History", "Last season as a fraction of peak.", "Below 0.85 means meaningfully off his best."),
    ("fp_consistency", "History", "Standard deviation of fantasy points across five seasons.", "Low = predictable. High = boom/bust."),
    ("seasons_of_data", "History", "How many of the last five seasons have usable data.", "Slope on 2 seasons is fragile — treat with caution."),
    ("avg_gp_5y", "History", "Average games played per season over five years.", "The durability track record behind this year's number."),
    ("steal", "Pricing", "", ""),
    ("fp_per_36", "Rate", "League-scored fantasy points per 36 minutes.", "Role-independent rate. High value + low minutes = breakout candidate."),
    ("stocks_pg", "Rate", "Steals plus blocks per game.", "The single most underpriced stat in this format."),
    ("stock_fpts", "Rate", "Fantasy points from steals and blocks alone (3.0 each).", "Shows how much value is defensive."),
    ("stock_share", "Rate", "Fraction of total production coming from stocks.", "High share = value the wider market systematically misses."),
    ("ast_fpts", "Rate", "Fantasy points from assists (1.5 each).", ""),
    ("reb_fpts", "Rate", "Fantasy points from rebounds (1.2 each).", ""),
    ("tov_drag", "Rate", "Fantasy points lost to turnovers.", "Modest penalty here; do not overweight it."),
    ("custom_score_pg", "Production", "Last season's fantasy points per game under THIS league's scoring.", "The core production number."),
    ("adj_score", "Production", "Fantasy points per game after the availability discount.", "What actually drives the ranking."),
    ("durability", "Production", "Availability multiplier from games played (0.70-1.00).", "Weekly H2H pays nothing for games missed."),
    ("GP", "Production", "Games played last season.", ""),
    ("MIN", "Production", "Minutes per game last season.", "Rising minutes often precede a breakout."),
    ("PTS", "Production", "Points per game.", ""),
    ("REB", "Production", "Rebounds per game.", ""),
    ("AST", "Production", "Assists per game.", ""),
    ("STL", "Production", "Steals per game.", ""),
    ("BLK", "Production", "Blocks per game.", ""),
    ("TOV", "Production", "Turnovers per game.", ""),
    ("yoy_trend", "Production", "Change in fantasy points vs the prior season.", ""),
    ("late_trend", "Production", "Final-20-game form vs full-season average.", "Negative means he finished cold."),
    ("min_trend", "Production", "Change in minutes over the final 20 games.", "The clearest early signal of a role change."),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
GROUP_FILL = PatternFill("solid", fgColor="D9E2F3")

def _autosize(ws, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(longest + 2, 9), max_width)

def _style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

def write_draft_workbook(df, path, scarcity_text="", meta=None):
    """Writes the draft plan as a multi-tab workbook: the plan itself, a column guide,
    quick-reference target/fade views, and the run's provenance."""
    present = [c for c in df.columns]
    guide_rows, seen = [], set()
    for col, group, meaning, use in COLUMN_GUIDE:
        if col in present and col not in seen and meaning:
            seen.add(col)
            guide_rows.append({"Column": col, "Group": group, "What it means": meaning,
                               "How to use it": use})
    # Surface any column that exists but was never documented.
    for col in present:
        if col not in seen:
            guide_rows.append({"Column": col, "Group": "Other", "What it means": "", "How to use it": ""})
    guide = pd.DataFrame(guide_rows)

    targets = df[(df.get("market_vs_value") == "underrated")].head(60)
    fades = df[(df.get("market_vs_value") == "overhyped")].head(60)
    short_cols = [c for c in ["rank", "player", "position", "team", "tier", "trend",
                              "steal", "target", "walk_away", "mechanical_value",
                              "career_arc", "why", "live_note"] if c in df.columns]

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name="Draft Plan", index=False)
        guide.to_excel(xl, sheet_name="Column Guide", index=False)
        targets[short_cols].to_excel(xl, sheet_name="Bargain Targets", index=False)
        fades[short_cols].to_excel(xl, sheet_name="Fades", index=False)

        info = [("Generated", meta.get("generated", "") if meta else "")]
        if meta:
            info += [(k, v) for k, v in meta.items() if k != "generated"]
        if scarcity_text:
            info += [("", "")] + [(line.strip(), "") for line in scarcity_text.split("\n")]
        pd.DataFrame(info, columns=["Item", "Value"]).to_excel(
            xl, sheet_name="Run Info", index=False
        )

        for name in ("Draft Plan", "Column Guide", "Bargain Targets", "Fades", "Run Info"):
            ws = xl.book[name]
            _style_header(ws)
            _autosize(ws)

        # Long prose columns need a fixed width or they blow out the sheet.
        plan = xl.book["Draft Plan"]
        for col_name in ("community_read", "why", "risk", "live_note", "trajectory"):
            if col_name in present:
                letter = get_column_letter(present.index(col_name) + 1)
                plan.column_dimensions[letter].width = 60
        guide_ws = xl.book["Column Guide"]
        guide_ws.column_dimensions["C"].width = 62
        guide_ws.column_dimensions["D"].width = 62
    return path
